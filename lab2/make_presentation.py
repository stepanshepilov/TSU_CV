from pathlib import Path

import cv2
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
IMAGE_PATH = ROOT / "image.png"
OUTPUT_PATH = ROOT / "Лабораторная_2_вариант_9.pptx"
ASSET_DIR = ROOT / ".presentation_assets"
ASSET_DIR.mkdir(exist_ok=True)

image = cv2.imread(str(IMAGE_PATH), cv2.IMREAD_GRAYSCALE)
laplacian = np.abs(cv2.Laplacian(image, cv2.CV_64F))
histogram = cv2.calcHist([image], [0], None, [256], [0, 256]).ravel()
otsu_threshold, binary = cv2.threshold(image, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
horizontal_projection = image.sum(axis=1)
vertical_projection = image.sum(axis=0)
lap_horizontal = laplacian.mean(axis=1)
lap_vertical = laplacian.mean(axis=0)
binary_horizontal = binary.sum(axis=1)
binary_vertical = binary.sum(axis=0)


def save_plot(name, plotter, figsize=(10, 4)):
    path = ASSET_DIR / name
    plt.figure(figsize=figsize)
    plotter()
    plt.tight_layout()
    plt.savefig(path, dpi=180, bbox_inches="tight")
    plt.close()
    return path


hist_path = save_plot(
    "hist.png",
    lambda: (plt.plot(histogram, color="#162536"), plt.axvline(np.argmax(histogram), color="#e45c4f", linestyle="--"), plt.title("Гистограмма исходного изображения"), plt.xlabel("Интенсивность"), plt.ylabel("Количество пикселей")),
)
lap_path = save_plot(
    "laplacian.png",
    lambda: (plt.plot(lap_horizontal, label="по строкам", color="#0e7490"), plt.plot(lap_vertical, label="по столбцам", color="#e45c4f"), plt.title("Профили модуля отклика Лапласа"), plt.xlabel("Координата"), plt.ylabel("Средний |отклик|"), plt.legend()),
)
projection_path = save_plot(
    "projections.png",
    lambda: (plt.plot(horizontal_projection, label="горизонтальная", color="#0e7490"), plt.plot(vertical_projection, label="вертикальная", color="#e45c4f"), plt.title("Проекции исходного изображения"), plt.xlabel("Координата"), plt.ylabel("Сумма интенсивностей"), plt.legend()),
)
binary_path = save_plot(
    "binary.png",
    lambda: (plt.imshow(binary, cmap="gray"), plt.title(f"Бинаризация Отсу, порог = {otsu_threshold:.0f}"), plt.axis("off")),
    figsize=(8, 4),
)
binary_projection_path = save_plot(
    "binary_projections.png",
    lambda: (plt.plot(binary_horizontal, label="по строкам", color="#0e7490"), plt.plot(binary_vertical, label="по столбцам", color="#e45c4f"), plt.title("Проекции после бинаризации Отсу"), plt.xlabel("Координата"), plt.ylabel("Сумма белых пикселей"), plt.legend()),
)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = RGBColor(246, 243, 237)
INK = RGBColor(22, 37, 54)
TEAL = RGBColor(14, 116, 144)
CORAL = RGBColor(228, 92, 79)
MUTED = RGBColor(84, 99, 108)


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.alignment = align
        paragraph.space_after = Pt(7)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.45, 12, 0.65, size=29, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.12, 11.8, 0.4, size=13, color=MUTED)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(1.62), Inches(1.0), Inches(0.08))
    rule.fill.solid()
    rule.fill.fore_color.rgb = CORAL
    rule.line.fill.background()


def add_footer(slide, number):
    add_text(slide, f"Лабораторная 2 · вариант 9                                      {number:02d}", 0.68, 7.12, 12, 0.2, size=9, color=MUTED)


def new_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_text(slide, "ЛАБОРАТОРНАЯ 02", 0.75, 0.75, 6, 0.4, size=15, color=TEAL, bold=True)
add_text(slide, "Гистограммы, профили\nи проекции изображения", 0.72, 1.45, 7.4, 1.7, size=34, bold=True)
add_text(slide, "Компьютерное зрение · вариант 9", 0.78, 3.45, 5.5, 0.5, size=18, color=MUTED)
slide.shapes.add_picture(str(IMAGE_PATH), Inches(8.2), Inches(0.6), width=Inches(4.45), height=Inches(2.5))
add_text(slide, "Практика: измеряем яркость, границы и положение объекта через одномерные сигналы", 0.78, 5.65, 8.5, 0.6, size=16, color=INK)
add_footer(slide, 1)

slide = new_slide("Задание варианта 9", "Пять операций над одним изображением")
add_text(slide, "01  Гистограмма серого изображения и её пик\n02  Профили после фильтра Лапласа\n03  Горизонтальная проекция и строка максимума\n04  Вертикальная проекция и столбец максимума\n05  Повтор профилей и проекций после порога Отсу", 0.9, 2.0, 8.0, 3.5, size=22)
add_text(slide, "Вход: image.png\nРазмер: 1000 × 563\nВыход: графики + численные координаты экстремумов", 9.2, 2.2, 3.2, 2.2, size=17, color=TEAL, bold=True)

slide = new_slide("Теория: три одномерных взгляда", "Сворачиваем двумерную матрицу яркости в измеримые сигналы")
add_text(slide, "Гистограмма\nH(k) = число пикселей с яркостью k\n\nПрофиль\nP(y) = (1/N) Σₓ I(x, y)\nP(x) = (1/M) Σᵧ I(x, y)\n\nПроекция\nHP(y) = Σₓ I(x, y)\nVP(x) = Σᵧ I(x, y)", 0.85, 1.95, 5.1, 4.5, size=21)
add_text(slide, "Гистограмма отвечает: «какие тона преобладают?»\n\nПрофиль отвечает: «какова средняя яркость в каждой строке или колонке?»\n\nПроекция отвечает: «где накоплено больше суммарной яркости?»", 6.6, 2.05, 5.8, 3.8, size=20, color=MUTED)

slide = new_slide("Исходная гистограмма", "Пик показывает доминирующий уровень яркости")
slide.shapes.add_picture(str(hist_path), Inches(0.75), Inches(1.85), width=Inches(7.9), height=Inches(4.65))
add_text(slide, "Пик: интенсивность 225\nКоличество: 40 765 пикселей", 9.1, 2.5, 3.2, 1.5, size=24, color=CORAL, bold=True)
add_text(slide, "Гистограмма не сообщает координаты пикселей. Она описывает распределение яркостей по всему изображению.", 9.1, 4.35, 3.25, 1.2, size=16, color=MUTED)

slide = new_slide("Лаплас: где яркость меняется резко", "Вторая производная подчёркивает границы и мелкие детали")
add_text(slide, "laplacian_signed = cv2.Laplacian(image, cv2.CV_64F)\nlaplacian = np.abs(laplacian_signed)\n\nprofile_y = laplacian.mean(axis=1)\nprofile_x = laplacian.mean(axis=0)", 0.85, 2.0, 5.1, 2.8, size=18, color=INK)
slide.shapes.add_picture(str(lap_path), Inches(6.25), Inches(1.85), width=Inches(6.3), height=Inches(4.3))
add_text(slide, "Максимум по строкам: y = 562\nМаксимум по столбцам: x = 535", 0.9, 5.35, 4.8, 0.8, size=19, color=TEAL, bold=True)

slide = new_slide("Проекции исходного изображения", "Сумма яркостей вдоль каждой координаты")
slide.shapes.add_picture(str(projection_path), Inches(0.65), Inches(1.8), width=Inches(8.0), height=Inches(4.7))
add_text(slide, "Горизонтальная\nстрока y = 86\nсумма = 223 754\n\nВертикальная\nстолбец x = 372\nсумма = 129 309", 9.1, 2.05, 3.1, 3.9, size=21, color=TEAL, bold=True)

slide = new_slide("Отсу: превращаем яркость в классы", "Автоматический порог через максимум межклассовой дисперсии")
slide.shapes.add_picture(str(binary_path), Inches(0.75), Inches(1.9), width=Inches(6.7), height=Inches(4.6))
add_text(slide, "otsu_threshold, binary = cv2.threshold(\n    image, 0, 255,\n    cv2.THRESH_BINARY + cv2.THRESH_OTSU\n)", 7.8, 2.0, 4.6, 1.5, size=17)
add_text(slide, "Порог: 141\nБелые пиксели: 85.04%\n\nПосле порога проекции можно читать как число белых пикселей в строке или столбце.", 7.8, 4.15, 4.3, 1.8, size=20, color=CORAL, bold=True)

slide = new_slide("Профили и проекции после Отсу", "Бинаризация делает геометрию объекта более заметной")
slide.shapes.add_picture(str(binary_projection_path), Inches(0.75), Inches(1.85), width=Inches(8.1), height=Inches(4.7))
add_text(slide, "Для бинарного изображения:\n\nпрофиль = средняя доля белого\n\nпроекция / 255 = количество белых пикселей", 9.2, 2.1, 3.2, 2.6, size=20, color=TEAL, bold=True)

slide = new_slide("Код лабораторной: общий конвейер", "Одна и та же последовательность повторяется для исходного и бинарного изображения")
add_text(slide, "image = cv2.imread('image.png', cv2.IMREAD_GRAYSCALE)\n\nhist = cv2.calcHist([image], [0], None, [256], [0, 256])\n\nhorizontal_profile = image.mean(axis=1)\nvertical_profile = image.mean(axis=0)\n\nhorizontal_projection = image.sum(axis=1)\nvertical_projection = image.sum(axis=0)\n\nthreshold, binary = cv2.threshold(image, 0, 255,\n                                  cv2.THRESH_BINARY + cv2.THRESH_OTSU)", 0.85, 1.85, 11.7, 4.8, size=18, color=INK)

slide = new_slide("Как читать результаты", "Разные операции отвечают на разные вопросы")
add_text(slide, "Гистограмма\n225 — самый частый тон исходного изображения.\n\nЛаплас\nВысокий профиль — сильная локальная граница или деталь.\n\nПроекция\nПик показывает координату с наибольшей суммарной яркостью.\n\nОтсу\nПорог 141 разделяет изображение на два класса, после чего пики проекций отражают геометрию белой области.", 0.95, 1.9, 11.2, 4.7, size=22)

slide = new_slide("Контрольные вопросы 1–5", "Короткие ответы по базовым определениям")
add_text(slide, "1. Гистограмма — распределение числа пикселей по яркостям; нужна для анализа тонов и контраста.\n\n2. У серого изображения один канал, у цветного обычно отдельная гистограмма для каждого канала.\n\n3. Профиль — одномерный массив средних интенсивностей по строкам или столбцам.\n\n4. Горизонтальный профиль: mean(axis=1), вертикальный: mean(axis=0).\n\n5. Проекция — сумма интенсивностей вдоль одной координаты.", 0.85, 1.85, 11.6, 4.95, size=19)

slide = new_slide("Контрольные вопросы 6–10", "Ответы, связанные с интерпретацией")
add_text(slide, "6. Профиль использует среднее, проекция — сумму; поэтому проекция зависит от размера строки или столбца.\n\n7. По гистограмме можно выбрать растяжение контраста или эквализацию.\n\n8. Бинарные проекции находят строки текста, границы, положение и размеры объектов.\n\n9. Сглаживание уменьшает шум и обычно делает гистограмму компактнее и менее зубчатой.\n\n10. Пик — самый частый уровень яркости и число пикселей на нём.", 0.85, 1.85, 11.6, 4.95, size=19)

slide = new_slide("Итоги", "Что выполнено в ноутбуке")
add_text(slide, "✓ Найден пик исходной гистограммы: 225\n✓ Построены профили модуля отклика Лапласа\n✓ Найдены максимумы горизонтальной и вертикальной проекций\n✓ Выполнена бинаризация Отсу с порогом 141\n✓ Построены профили и проекции бинарного изображения\n✓ Даны ответы на 10 контрольных вопросов", 1.0, 2.05, 9.0, 3.5, size=25, color=INK)
add_text(slide, "Главная идея: двумерное изображение можно исследовать через гистограмму и одномерные сигналы, сохраняя полезную информацию о яркости, границах и положении объекта.", 1.0, 5.6, 11.1, 0.7, size=18, color=TEAL, bold=True)

prs.save(OUTPUT_PATH)
print(f"saved: {OUTPUT_PATH}")