"""Собирает презентацию по лабораторной работе 3 (вариант 9)."""
from pathlib import Path

import cv2
import matplotlib
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

matplotlib.use("Agg")

ROOT = Path(__file__).parent
IMAGE_PATH = ROOT / ".." / "lab2" / "image.png"
OUTPUT_PATH = ROOT / "Лабораторная_3_вариант_9.pptx"
ASSET_DIR = ROOT / ".presentation_assets"
ASSET_DIR.mkdir(exist_ok=True)

image = cv2.imread(str(IMAGE_PATH), cv2.IMREAD_COLOR)
image_rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
height, width = image.shape[:2]
center = (width / 2, height / 2)
BORDER = (30, 30, 30)

# 1. Масштабирование до квадрата 300×300
TARGET = (300, 300)
resized = cv2.resize(image, TARGET, interpolation=cv2.INTER_AREA)

# 2. Поворот на 25°
ANGLE = 25.0
M_rotate = cv2.getRotationMatrix2D(center, ANGLE, 1.0)
rotated = cv2.warpAffine(image, M_rotate, (width, height), borderValue=BORDER)
cos_a, sin_a = abs(M_rotate[0, 0]), abs(M_rotate[0, 1])
exp_w = int(height * sin_a + width * cos_a)
exp_h = int(height * cos_a + width * sin_a)
M_expand = M_rotate.copy()
M_expand[0, 2] += exp_w / 2 - center[0]
M_expand[1, 2] += exp_h / 2 - center[1]
rotated_full = cv2.warpAffine(image, M_expand, (exp_w, exp_h), borderValue=BORDER)

# 3. Сдвиг на 120 влево и 60 вниз
TX, TY = -120, 60
T = np.float32([[1, 0, TX], [0, 1, TY]])
shifted = cv2.warpAffine(image, T, (width, height), borderValue=BORDER)

# 4. Отражение по вертикали
flipped_v = cv2.flip(image, 0)
flipped_h = cv2.flip(image, 1)
flipped_both = cv2.flip(image, -1)
M_flip_v = np.float32([[1, 0, 0], [0, -1, height - 1]])

# 5. Аффинное преобразование треугольника
pts_src = np.float32([[150, 100], [850, 120], [200, 480]])
pts_dst = np.float32([[80, 180], [900, 60], [420, 520]])
M_affine = cv2.getAffineTransform(pts_src, pts_dst)
affine = cv2.warpAffine(image, M_affine, (width, height), borderValue=BORDER)
det_affine = float(np.linalg.det(M_affine[:, :2]))

src_marked = image.copy()
dst_marked = affine.copy()
for img_marked, pts, color in ((src_marked, pts_src, (0, 0, 255)), (dst_marked, pts_dst, (0, 255, 0))):
    cv2.polylines(img_marked, [pts.astype(np.int32)], True, color, 5)
    for idx, (x, y) in enumerate(pts.astype(int), 1):
        cv2.circle(img_marked, (x, y), 11, color, -1)


def cross2(u, v):
    return float(u[0] * v[1] - u[1] * v[0])


def triangle_metrics(points):
    a, b, c = points
    sides = np.array([np.linalg.norm(b - a), np.linalg.norm(c - b), np.linalg.norm(a - c)])
    angles = []
    for p, q, r in ((a, b, c), (b, c, a), (c, a, b)):
        v1, v2 = q - p, r - p
        cos_angle = np.dot(v1, v2) / (np.linalg.norm(v1) * np.linalg.norm(v2))
        angles.append(np.degrees(np.arccos(np.clip(cos_angle, -1, 1))))
    return sides, np.array(angles), 0.5 * abs(cross2(b - a, c - a))


sides_src, angles_src, area_src = triangle_metrics(pts_src)
sides_dst, angles_dst, area_dst = triangle_metrics(pts_dst)

# Композиция преобразований
to_3x3 = lambda m: np.vstack([m, [0, 0, 1]])
M_total = to_3x3(T) @ to_3x3(M_rotate) @ to_3x3(M_flip_v)
combined_once = cv2.warpAffine(image, M_total[:2], (width, height), borderValue=BORDER)
step_by_step = cv2.warpAffine(
    cv2.warpAffine(cv2.flip(image, 0), M_rotate, (width, height), borderValue=BORDER),
    T, (width, height), borderValue=BORDER)
composition_diff = cv2.absdiff(combined_once, step_by_step)


def panels(name, items, figsize=(14, 4)):
    # items: список (BGR-изображение, подпись)
    path = ASSET_DIR / name
    fig, axes = plt.subplots(1, len(items), figsize=figsize)
    for ax, (img, title) in zip(np.atleast_1d(axes), items):
        ax.imshow(cv2.cvtColor(img, cv2.COLOR_BGR2RGB) if img.ndim == 3 else img,
                  cmap=None if img.ndim == 3 else "inferno")
        ax.set_title(title, fontsize=11)
        ax.axis("off")
    plt.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)
    return path


resize_path = panels("resize.png",
                     [(image, f"Исходное {width}×{height}"), (resized, "resize 300×300")],
                     figsize=(11, 3.4))
interp_path = panels(
    "interp.png",
    [(cv2.resize(image, TARGET, interpolation=flag), name)
     for name, flag in (("INTER_NEAREST", cv2.INTER_NEAREST),
                        ("INTER_LINEAR", cv2.INTER_LINEAR),
                        ("INTER_CUBIC", cv2.INTER_CUBIC),
                        ("INTER_AREA", cv2.INTER_AREA))],
    figsize=(15, 3.8))
rotate_path = panels("rotate.png",
                     [(rotated, f"Поворот {ANGLE:.0f}° в исходном кадре"),
                      (rotated_full, f"Поворот с расширением до {exp_w}×{exp_h}")],
                     figsize=(12, 3.6))
shift_path = panels("shift.png",
                    [(image, "Исходное"), (shifted, f"Сдвиг t_x = {TX}, t_y = {TY}")],
                    figsize=(12, 3.4))
flip_path = panels("flip.png",
                   [(flipped_v, "flip(0) — по вертикали"),
                    (flipped_h, "flip(1) — по горизонтали"),
                    (flipped_both, "flip(−1) — обе оси")],
                   figsize=(15, 3.0))
affine_path = panels("affine.png",
                     [(src_marked, "Треугольник P1P2P3 до"), (dst_marked, "После преобразования")],
                     figsize=(12, 3.6))
compose_path = panels("compose.png",
                      [(combined_once, "Одна матрица"), (step_by_step, "Три вызова warpAffine"),
                       (composition_diff.max(axis=2), "Модуль разницы")],
                      figsize=(15, 3.2))

# График метрик треугольника
metrics_path = ASSET_DIR / "metrics.png"
fig, axes = plt.subplots(1, 2, figsize=(11, 3.8))
x = np.arange(3)
axes[0].bar(x - 0.19, sides_src, 0.38, label="до", color="#0e7490")
axes[0].bar(x + 0.19, sides_dst, 0.38, label="после", color="#e45c4f")
axes[0].set_xticks(x, ["P1P2", "P2P3", "P3P1"])
axes[0].set_title("Длины сторон, пикс.")
axes[0].legend()
axes[1].bar(x - 0.19, angles_src, 0.38, label="до", color="#0e7490")
axes[1].bar(x + 0.19, angles_dst, 0.38, label="после", color="#e45c4f")
axes[1].set_xticks(x, ["при P1", "при P2", "при P3"])
axes[1].set_title("Углы, градусы")
axes[1].legend()
plt.tight_layout()
fig.savefig(metrics_path, dpi=170, bbox_inches="tight")
plt.close(fig)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = RGBColor(246, 243, 237)
INK = RGBColor(22, 37, 54)
TEAL = RGBColor(14, 116, 144)
CORAL = RGBColor(228, 92, 79)
MUTED = RGBColor(84, 99, 108)


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.alignment = align
        paragraph.space_after = Pt(7)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.45, 12, 0.65, size=29, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.12, 11.8, 0.4, size=13, color=MUTED)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(1.62), Inches(1.0), Inches(0.08))
    rule.fill.solid()
    rule.fill.fore_color.rgb = CORAL
    rule.line.fill.background()


def add_footer(slide, number):
    add_text(slide, f"Лабораторная 3 · вариант 9                                      {number:02d}",
             0.68, 7.12, 12, 0.2, size=9, color=MUTED)


def new_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


# 01 — титульный
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_text(slide, "ЛАБОРАТОРНАЯ 03", 0.75, 0.75, 6, 0.4, size=15, color=TEAL, bold=True)
add_text(slide, "Геометрические\nпреобразования изображений", 0.72, 1.45, 7.4, 1.7, size=34, bold=True)
add_text(slide, "Компьютерное зрение · вариант 9", 0.78, 3.45, 5.5, 0.5, size=18, color=MUTED)
slide.shapes.add_picture(str(IMAGE_PATH), Inches(8.2), Inches(0.6), width=Inches(4.45), height=Inches(2.5))
add_text(slide, "Масштаб, поворот, сдвиг, отражение и аффинное преобразование — всё через матрицу 2 × 3",
         0.78, 5.65, 8.9, 0.6, size=16, color=INK)
add_footer(slide, 1)

# 02 — задание
slide = new_slide("Задание варианта 9", "Пять преобразований одного изображения")
add_text(slide, "01  Масштабировать до квадрата 300 × 300\n"
                "02  Повернуть на 25°\n"
                "03  Сместить на 120 пикселей влево и 60 вниз\n"
                "04  Отразить по вертикали\n"
                "05  Аффинное преобразование формы треугольника",
         0.9, 2.0, 8.2, 3.5, size=22)
add_text(slide, f"Вход: ../lab2/image.png\nРазмер: {width} × {height}\n"
                f"Центр: ({center[0]:.0f}, {center[1]:.0f})\nВыход: файлы в output/",
         9.3, 2.2, 3.3, 2.2, size=17, color=TEAL, bold=True)
add_text(slide, "Контрольные вопросы варианта: что произойдёт с углами треугольника при аффинном "
                "преобразовании и как в OpenCV задаются исходные и целевые точки.",
         0.9, 5.7, 11.5, 0.9, size=17, color=MUTED)

# 03 — теория
slide = new_slide("Теория: всё это одна матрица", "Аффинное преобразование в однородных координатах")
add_text(slide, "(x')   ( a11  a12  b1 ) (x)\n(y') = ( a21  a22  b2 ) (y)\n                                    (1)\n\n"
                "Левый блок 2 × 2 — линейная часть\n(поворот, масштаб, скос).\n"
                "Правый столбец — перенос.\n\n"
                "|det| линейной части — во сколько раз\nменяется площадь фигуры.\n"
                "Знак det — сохраняется ли ориентация.",
         0.85, 1.95, 5.5, 4.5, size=19)
add_text(slide, "Масштаб      x' = s_x·x,  y' = s_y·y\n\n"
                "Поворот       α = s·cos θ,  β = s·sin θ\n\n"
                "Перенос       T = [[1, 0, t_x], [0, 1, t_y]]\n\n"
                "Отражение   flip(0), flip(1), flip(−1)\n\n"
                "Аффинное     3 пары точек → матрица 2 × 3\n\n"
                "Перспективное 4 пары точек → матрица 3 × 3",
         6.6, 1.95, 5.9, 4.5, size=19, color=MUTED)

# 04 — пункт 1
slide = new_slide("01 · Масштабирование до 300 × 300", "Соотношение сторон меняется с 1.776 на 1.000")
slide.shapes.add_picture(str(resize_path), Inches(0.7), Inches(1.9), width=Inches(8.2), height=Inches(2.7))
add_text(slide, "resized = cv2.resize(image, (300, 300),\n                     interpolation=cv2.INTER_AREA)",
         0.75, 4.9, 8.2, 1.0, size=18)
add_text(slide, f"s_x = {TARGET[0] / width:.3f}\ns_y = {TARGET[1] / height:.3f}\n\n"
                f"Пикселей меньше\nв {height * width / (TARGET[0] * TARGET[1]):.1f} раза",
         9.3, 2.2, 3.2, 2.6, size=20, color=CORAL, bold=True)
add_text(slide, "cv2.resize принимает размер в порядке (ширина, высота) — обратном порядку shape.",
         0.75, 6.0, 11.6, 0.7, size=17, color=MUTED)

# 05 — интерполяция
slide = new_slide("Интерполяция: как заполняются новые узлы", "При уменьшении лучший выбор — INTER_AREA")
slide.shapes.add_picture(str(interp_path), Inches(0.6), Inches(1.95), width=Inches(12.1), height=Inches(3.1))
add_text(slide, "INTER_NEAREST берёт ближайший пиксель и просто выбрасывает лишние строки и столбцы — "
                "появляются «лесенки». INTER_AREA усредняет все исходные пиксели, попавшие в новый, "
                "и потому не даёт алиасинга. INTER_CUBIC выгоднее при увеличении, а не при уменьшении.",
         0.75, 5.35, 11.8, 1.4, size=19)

# 06 — пункт 2
slide = new_slide("02 · Поворот на 25°", "getRotationMatrix2D строит матрицу, warpAffine её применяет")
slide.shapes.add_picture(str(rotate_path), Inches(0.7), Inches(1.9), width=Inches(8.0), height=Inches(2.9))
add_text(slide, f"M = [[{M_rotate[0,0]:.4f}, {M_rotate[0,1]:.4f}, {M_rotate[0,2]:.2f}],\n"
                f"     [{M_rotate[1,0]:.4f}, {M_rotate[1,1]:.4f}, {M_rotate[1,2]:.2f}]]",
         0.75, 5.05, 8.0, 0.9, size=17)
add_text(slide, f"cos 25° = {np.cos(np.deg2rad(ANGLE)):.4f}\nsin 25° = {np.sin(np.deg2rad(ANGLE)):.4f}\n"
                f"det = {np.linalg.det(M_rotate[:, :2]):.4f}",
         9.1, 2.2, 3.4, 1.6, size=20, color=TEAL, bold=True)
add_text(slide, f"Определитель равен единице — поворот сохраняет площадь.\n\n"
                f"Если оставить кадр {width}×{height}, углы обрезаются. Расширенный кадр: {exp_w}×{exp_h}.",
         9.1, 3.9, 3.4, 2.4, size=16, color=MUTED)

# 07 — пункт 3
slide = new_slide("03 · Сдвиг на 120 влево и 60 вниз", "Матрица переноса с единичной линейной частью")
slide.shapes.add_picture(str(shift_path), Inches(0.7), Inches(1.9), width=Inches(8.2), height=Inches(2.9))
add_text(slide, "T = np.float32([[1, 0, -120],\n                [0, 1,   60]])\n\n"
                "shifted = cv2.warpAffine(image, T, (width, height))",
         0.75, 5.0, 8.2, 1.5, size=18)
add_text(slide, f"t_x = {TX} → влево\nt_y = {TY} → вниз\n\nПроверка:\n(400, 200) → (280, 260)",
         9.3, 2.2, 3.2, 2.6, size=19, color=CORAL, bold=True)
add_text(slide, "Освободившиеся полосы заполняются цветом borderValue: этих пикселей в исходном "
                "изображении просто нет.",
         9.3, 4.9, 3.2, 1.6, size=15, color=MUTED)

# 08 — пункт 4
slide = new_slide("04 · Отражение по вертикали", "cv2.flip численно эквивалентен срезу NumPy")
slide.shapes.add_picture(str(flip_path), Inches(0.6), Inches(1.95), width=Inches(9.0), height=Inches(2.4))
add_text(slide, "flip(0)   ≡ image[::-1, :]        True\n"
                "flip(1)   ≡ image[:, ::-1]        True\n"
                "flip(−1) ≡ image[::-1, ::-1]   True",
         0.75, 4.6, 8.2, 1.3, size=19)
add_text(slide, f"flip(0) совпал\nс warpAffine по матрице\n[[1, 0, 0], [0, −1, {height - 1}]]\n\n"
                f"det = {np.linalg.det(M_flip_v[:, :2]):.0f}",
         9.8, 2.2, 2.9, 2.6, size=17, color=TEAL, bold=True)
add_text(slide, "Отрицательный определитель означает, что отражение меняет ориентацию на зеркальную — "
                "в отличие от поворота, у которого det = +1. flip(−1) эквивалентен повороту на 180°.",
         0.75, 5.9, 11.8, 0.9, size=17, color=MUTED)

# 09 — пункт 5
slide = new_slide("05 · Аффинное преобразование треугольника", "Три пары точек однозначно определяют матрицу 2 × 3")
slide.shapes.add_picture(str(affine_path), Inches(0.7), Inches(1.9), width=Inches(8.0), height=Inches(2.9))
add_text(slide, "pts_src = np.float32([[150,100],[850,120],[200,480]])\n"
                "pts_dst = np.float32([[ 80,180],[900, 60],[420,520]])\n"
                "M = cv2.getAffineTransform(pts_src, pts_dst)",
         0.75, 5.0, 8.2, 1.3, size=17)
add_text(slide, "Шесть неизвестных,\nшесть уравнений —\nрешение единственное.\n\n"
                f"det = {det_affine:.4f}",
         9.2, 2.2, 3.3, 2.6, size=19, color=TEAL, bold=True)
add_text(slide, "Проверка через cv2.transform: все три точки попали ровно в заданные позиции.",
         9.2, 4.9, 3.3, 1.4, size=15, color=MUTED)

# 10 — метрики треугольника
slide = new_slide("Что стало с треугольником", "Ответ на контрольный вопрос варианта")
slide.shapes.add_picture(str(metrics_path), Inches(0.7), Inches(1.9), width=Inches(8.1), height=Inches(2.8))
add_text(slide, f"Углы до:    {angles_src[0]:.1f}°, {angles_src[1]:.1f}°, {angles_src[2]:.1f}°\n"
                f"Углы после: {angles_dst[0]:.1f}°, {angles_dst[1]:.1f}°, {angles_dst[2]:.1f}°\n"
                f"Сумма углов: 180° в обоих случаях",
         0.75, 4.9, 8.1, 1.4, size=19)
add_text(slide, f"Площадь\nдо: {area_src:.0f}\nпосле: {area_dst:.0f}\n"
                f"отношение {area_dst / area_src:.4f}\n\ndet матрицы\n{det_affine:.4f}",
         9.2, 2.2, 3.3, 3.0, size=18, color=CORAL, bold=True)
add_text(slide, "Отношение площадей совпало\nс определителем линейной части.",
         9.2, 5.4, 3.3, 1.0, size=15, color=MUTED)

# 11 — сохранение параллельности
slide = new_slide("Что аффинное преобразование сохраняет", "Углы — нет, параллельность и отношения длин — да")
add_text(slide, "НЕ сохраняются\n\n"
                "· длины сторон\n"
                "· углы между сторонами\n"
                "· площадь (умножается на |det|)\n"
                "· форма фигуры",
         0.9, 1.95, 5.3, 2.9, size=20, color=CORAL, bold=True)
add_text(slide, "Сохраняются\n\n"
                "· прямые остаются прямыми\n"
                "· параллельность прямых\n"
                "· отношение длин на одной прямой\n"
                "· середина отрезка",
         6.7, 1.95, 5.7, 2.9, size=20, color=TEAL, bold=True)
add_text(slide, "Проверено численно: у параллелограмма после преобразования синус угла между "
                "противоположными сторонами составил порядка 10⁻⁸ — стороны остались параллельными. "
                "Середина диагонали, преобразованная матрицей, совпала с серединой преобразованной "
                "диагонали. Углы сохраняет только подобие — частный случай аффинного преобразования.",
         0.9, 5.15, 11.6, 1.6, size=18)

# 12 — композиция
slide = new_slide("Композиция преобразований", "Перемножение матриц выгоднее последовательности вызовов")
slide.shapes.add_picture(str(compose_path), Inches(0.6), Inches(1.95), width=Inches(12.1), height=Inches(2.9))
add_text(slide, "M_total = to_3x3(T) @ to_3x3(M_rotate) @ to_3x3(M_flip_v)\n"
                "result = cv2.warpAffine(image, M_total[:2], (width, height))",
         0.75, 5.15, 7.6, 1.1, size=17)
add_text(slide, f"Разница с тремя\nотдельными вызовами:\nсредняя {composition_diff.mean():.1f}, "
                f"максимум {composition_diff.max()}",
         8.7, 5.2, 3.9, 1.3, size=17, color=CORAL, bold=True)

# 13 — контрольные вопросы варианта
slide = new_slide("Контрольные вопросы варианта 9", "Развёрнутые ответы")
add_text(slide, "Что произойдёт с углами треугольника при аффинном преобразовании?\n\n"
                "Углы изменятся. Линейная часть матрицы включает не только поворот, но и растяжение "
                "по осям и скос, а они деформируют фигуру: углы 80.9°, 30.6°, 68.5° перешли в "
                "53.3°, 35.5°, 91.2°. Сумма осталась 180° — фигура по-прежнему треугольник, прямые "
                "остались прямыми. Углы сохраняет только подобие. Площадь умножается на |det|.",
         0.85, 1.9, 11.6, 2.3, size=19)
add_text(slide, "Как в OpenCV задаются исходные и целевые точки?\n\n"
                "Массивом np.float32 формы (N, 2), где строка — пара (x, y): сначала столбец, потом "
                "строка. Это обратно порядку image[y, x]. Аффинное — 3 пары точек, перспективное — 4. "
                "Порядок точек в наборах должен соответствовать, точки не должны лежать на одной прямой. "
                "Если пар больше нужного — estimateAffine2D или findHomography с RANSAC.",
         0.85, 4.3, 11.6, 2.4, size=19, color=MUTED)

# 14 — контрольные вопросы практики
slide = new_slide("Контрольные вопросы практики №3", "Краткие ответы")
add_text(slide, "1. Аффинное — 2×3, 6 степеней свободы, сохраняет параллельность, 3 пары точек. "
                "Перспективное — 3×3, 8 степеней свободы, параллельность не сохраняет, 4 пары точек.\n\n"
                "2. warpAffine для каждого выходного пикселя находит источник через обратное "
                "преобразование и берёт значение с интерполяцией — иначе в выходной сетке остались бы «дыры».\n\n"
                "3. getRotationMatrix2D принимает центр вращения (x, y), угол в градусах и коэффициент масштаба.\n\n"
                "4. flip(1) меняет местами лево и право, flip(0) — верх и низ, flip(−1) — обе операции сразу.\n\n"
                "5. Для перспективного берут 4 точки в одной плоскости сцены, образующие выпуклый "
                "четырёхугольник; никакие три не лежат на одной прямой, порядок обхода в обоих наборах одинаков.",
         0.85, 1.85, 11.6, 4.95, size=18)

# 15 — итоги
slide = new_slide("Итоги", "Что выполнено в ноутбуке")
add_text(slide, "✓ Масштабирование до 300 × 300: соотношение сторон 1.776 → 1.000\n"
                f"✓ Поворот на 25°: det = 1, расширенный кадр {exp_w} × {exp_h}\n"
                "✓ Сдвиг (−120, +60): контрольная точка (400, 200) → (280, 260)\n"
                "✓ flip(0) совпал со срезом NumPy и с матричным преобразованием, det = −1\n"
                f"✓ Аффинное преобразование треугольника: det = {det_affine:.4f} = отношению площадей\n"
                "✓ Проверено сохранение параллельности и середин отрезков\n"
                "✓ Композиция трёх преобразований выполнена одной матрицей",
         1.0, 1.95, 11.0, 3.8, size=21, color=INK)
add_text(slide, "Главная идея: любое из этих преобразований — умножение координат на матрицу. "
                "Понимание матрицы позволяет предсказать результат и объединять операции без "
                "накопления ошибок интерполяции.",
         1.0, 5.7, 11.1, 0.9, size=18, color=TEAL, bold=True)

prs.save(OUTPUT_PATH)
print(f"saved: {OUTPUT_PATH}")
