from pathlib import Path

import cv2
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
IMAGE_PATH = ROOT.parent / 'lab2' / 'image.png'
OUTPUT = ROOT / 'Лабораторная_4_вариант_9.pptx'
ASSETS = ROOT / '.lab4_presentation_assets'
ASSETS.mkdir(exist_ok=True)
image = cv2.imread(str(IMAGE_PATH), cv2.IMREAD_GRAYSCALE)
blur5 = cv2.blur(image, (5, 5))
blur11 = cv2.blur(image, (11, 11))
rng = np.random.default_rng(42)
noisy = image.copy()
noise = rng.random(image.shape)
noisy[noise < .025] = 0
noisy[noise > .975] = 255
median = cv2.medianBlur(noisy, 5)
sobelx = cv2.Sobel(image, cv2.CV_64F, 1, 0, ksize=3)
sobely = cv2.Sobel(image, cv2.CV_64F, 0, 1, ksize=3)
sharpen = cv2.filter2D(image, -1, np.array([[0, -1, 0], [-1, 5, -1], [0, -1, 0]], np.float32))
spectrum = np.fft.fftshift(np.fft.fft2(image))
y, x = np.ogrid[:image.shape[0], :image.shape[1]]
distance = (x - image.shape[1] // 2) ** 2 + (y - image.shape[0] // 2) ** 2
hpf = 1 - np.exp(-distance / (2 * 45.0 ** 2))
high = np.abs(np.fft.ifft2(np.fft.ifftshift(spectrum * hpf)))
high = cv2.normalize(high, None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)


def plot(name, draw):
    path = ASSETS / name
    plt.figure(figsize=(9, 4))
    draw()
    plt.tight_layout()
    plt.savefig(path, dpi=160, bbox_inches='tight')
    plt.close()
    return path

compare = plot('compare.png', lambda: (plt.plot(image[281], label='исходное'), plt.plot(blur5[281], label='5x5'), plt.plot(blur11[281], label='11x11'), plt.title('Профиль строки после усреднения'), plt.legend()))
median_plot = plot('median.png', lambda: (plt.imshow(median, cmap='gray'), plt.title('Медианный фильтр 5x5'), plt.axis('off')))
sobel_plot = plot('sobel.png', lambda: (plt.imshow(np.sqrt(sobelx ** 2 + sobely ** 2), cmap='magma'), plt.title('Модуль градиента Собеля'), plt.axis('off')))
spectrum_plot = plot('spectrum.png', lambda: (plt.imshow(np.log1p(np.abs(spectrum)), cmap='magma'), plt.title('Логарифм спектра'), plt.axis('off')))
high_plot = plot('high.png', lambda: (plt.imshow(high, cmap='gray'), plt.title('Гауссов высокочастотный фильтр'), plt.axis('off')))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = RGBColor(246, 243, 237)
INK = RGBColor(22, 37, 54)
TEAL = RGBColor(14, 116, 144)
CORAL = RGBColor(228, 92, 79)
MUTED = RGBColor(84, 99, 108)


def add_text(slide, value, x, y, w, h, size=20, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.space_after = Pt(5)
    return box


def add_slide(title, body='', subtitle=''):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    add_text(slide, title, .65, .45, 12, .6, 29, INK, True)
    if subtitle:
        add_text(slide, subtitle, .68, 1.12, 12, .3, 13, MUTED)
    add_text(slide, body, .85, 1.9, 11.6, 4.9, 20, INK)
    add_text(slide, f'Лабораторная 4 · вариант 9                         {len(prs.slides):02d}', .68, 7.12, 12, .2, 9, MUTED)
    return slide


cover = add_slide('Цифровая фильтрация изображений', 'Компьютерное зрение · вариант 9')
cover.shapes.add_picture(str(IMAGE_PATH), Inches(8.2), Inches(.75), width=Inches(4.45), height=Inches(2.5))
add_text(cover, 'Сглаживание · шум · границы · резкость · БПФ', .85, 5.55, 8, .5, 18, TEAL, True)
add_slide('Задание варианта 9', '1. Усреднение 5x5 и 11x11\n2. Медианный фильтр против шума «соль и перец»\n3. Границы оператором Собеля\n4. Повышение резкости sharpen\n5. Гауссов ВЧ-фильтр в частотной области', 'Пять операций над image.png')
add_slide('Теория пространственной фильтрации', 'Свёртка вычисляет новый пиксель по его окрестности: g(x,y) = сумма h(i,j) f(x-i,y-j).\n\nЛинейные фильтры используют взвешенную сумму. Нелинейный медианный фильтр выбирает медиану окна. Большое окно сильнее сглаживает, но размывает детали.', 'Пространственная область')
s = add_slide('Усредняющие фильтры 5x5 и 11x11', 'Окно 5x5: среднее изменение 7.65\nОкно 11x11: среднее изменение 13.60', 'Чем больше окно, тем сильнее сглаживание')
s.shapes.add_picture(str(compare), Inches(.8), Inches(2.2), width=Inches(7.8), height=Inches(4.1))
add_text(s, '11x11 подавляет больше мелких изменений, но сильнее размывает границы.', 9, 2.7, 3.2, 1.2, 19, TEAL, True)
s = add_slide('Медианный фильтр', 'Добавлен воспроизводимый импульсный шум: ошибка 6.34. После medianBlur 5x5 ошибка стала 4.44.', 'Шум «соль и перец»')
s.shapes.add_picture(str(median_plot), Inches(.8), Inches(2.0), width=Inches(7), height=Inches(4.3))
add_text(s, 'Медиана удаляет одиночные выбросы и лучше сохраняет контуры, чем усреднение.', 8.25, 3, 4, 1.4, 20, CORAL, True)
s = add_slide('Оператор Собеля', 'Sobel X и Sobel Y оценивают первую производную яркости. Модуль градиента объединяет направления: G = sqrt(Gx^2 + Gy^2).', 'Выделение границ')
s.shapes.add_picture(str(sobel_plot), Inches(.8), Inches(2.0), width=Inches(7), height=Inches(4.3))
add_text(s, 'Максимум X: 1016\nМаксимум Y: 1016', 8.3, 3.2, 3.6, 1, 22, TEAL, True)
add_slide('Повышение резкости sharpen', 'Ядро [[0,-1,0],[-1,5,-1],[0,-1,0]] усиливает локальные перепады яркости. Среднее абсолютное изменение результата: 5.62.\n\nНедостаток: вместе с деталями может усилиться шум и появиться ореол около контуров.', 'Пространственное ядро')
s = add_slide('Частотная область и БПФ', 'Низкие частоты соответствуют плавному фону и крупным формам. Высокие частоты содержат границы, детали и шум.\n\nШаги: БПФ -> fftshift -> маска -> умножение спектров -> обратное БПФ.', 'Переход от пикселей к спектру')
s.shapes.add_picture(str(spectrum_plot), Inches(7.2), Inches(2.0), width=Inches(5.3), height=Inches(4.1))
s = add_slide('Гауссов высокочастотный фильтр', 'H_LPF(D) = exp(-D^2 / 2sigma^2)\nH_HPF(D) = 1 - H_LPF(D)\n\nsigma = 45. Низкие частоты подавляются плавно, высокие проходят и формируют карту контуров.', 'Фильтрация спектра')
s.shapes.add_picture(str(high_plot), Inches(7.2), Inches(2.0), width=Inches(5.3), height=Inches(4.1))
add_slide('Код лабораторной', "blur_5 = cv2.blur(image, (5, 5))\nblur_11 = cv2.blur(image, (11, 11))\nmedian = cv2.medianBlur(noisy, 5)\nsobel_x = cv2.Sobel(image, cv2.CV_64F, 1, 0, ksize=3)\nsharpened = cv2.filter2D(image, -1, kernel)\nspectrum = np.fft.fftshift(np.fft.fft2(image))\nfiltered = spectrum * gaussian_hpf", 'Основные вызовы из ноутбука')
add_slide('Контрольные вопросы 1–10', '1. Фильтрация нужна для шума, сглаживания, резкости и границ.\n2. Линейная использует сумму, нелинейная — медиану или другое правило.\n3. Шумы: гауссовский, импульсный, периодический и другие.\n4. Большое окно сильнее сглаживает и размывает детали.\n5. Пространственная удобна для локальных ядер, частотная — для диапазонов и больших ядер.\n6. Усреднение заменяет пиксель средним окна.\n7. Гаусс сильнее взвешивает центр.\n8. Медиана эффективна против «соли и перца».\n9. Собель и Превитт оценивают первую производную.\n10. Лаплас — вторая производная, Собель — направленная первая.', 'Ответы')
add_slide('Контрольные вопросы 11–20', '11. Sharpen усиливает разность с соседями или использует unsharp masking.\n12. Низкие частоты — фон, высокие — границы и детали.\n13. Переход выполняется дискретным преобразованием Фурье.\n14. БПФ, центрирование, маска, умножение, обратное БПФ.\n15. Идеальный НЧ имеет резкую границу, гауссов — плавную.\n16. Идеальный фильтр может дать ringing-ореолы.\n17. ВЧ-фильтр выделяет контуры.\n18. BPF пропускает диапазон частот.\n19. Частотный подход явно управляет диапазонами.\n20. БПФ выгоден при больших ядрах.', 'Ответы')
add_slide('Итоги', 'Выполнены все пункты варианта 9: сравнение окон, удаление импульсного шума, Собель, sharpen и гауссов ВЧ-фильтр через БПФ.\n\nГлавный вывод: размер окна и частотная маска задают компромисс между шумоподавлением, сохранением деталей и выделением границ.', 'Результат лабораторной работы')
prs.save(OUTPUT)
print(OUTPUT)
