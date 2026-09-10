from pathlib import Path
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT=Path(__file__).parent; OUT=ROOT/'Лабораторная_8_вариант_9.pptx'; ASSET=ROOT/'.presentation_assets'; ASSET.mkdir(exist_ok=True)
methods=['Порог 180','Отсу','SVM','U-Net']; scores=[.9505,1.0,.9995,.8504]
plt.figure(figsize=(8,4)); bars=plt.bar(methods,scores,color=['#d49a3a','#0e7490','#e45c4f','#162536']); plt.ylim(0,1.1); plt.ylabel('IoU относительно Отсу'); plt.title('Сравнение методов сегментации')
for b,v in zip(bars,scores): plt.text(b.get_x()+b.get_width()/2,v+.03,f'{v:.4f}',ha='center')
plt.tight_layout(); chart=ASSET/'comparison.png'; plt.savefig(chart,dpi=180); plt.close()
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG=RGBColor(246,243,237); INK=RGBColor(22,37,54); TEAL=RGBColor(14,116,144); CORAL=RGBColor(228,92,79); MUTED=RGBColor(84,99,108)
def text(s,v,x,y,w,h,size=20,color=INK,bold=False):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=b.text_frame; f.word_wrap=True; f.clear()
 for i,line in enumerate(v.split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph(); p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.color.rgb=color; p.font.bold=bold; p.space_after=Pt(5)
 return b
def slide(title,subtitle='',body='',size=20):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; text(s,title,.65,.45,12,.6,29,INK,True); text(s,subtitle,.68,1.12,12,.3,13,MUTED); text(s,body,.85,1.9,11.6,4.9,size=size); text(s,f'Лабораторная 8 · вариант 9                         {len(prs.slides):02d}',.68,7.12,12,.2,9,MUTED); return s
s=slide('Сегментация изображений','Вариант 9 · PyTorch'); s.shapes.add_picture(str(chart),Inches(7.8),Inches(1.1),width=Inches(4.8),height=Inches(2.5)); text(s,'Фиксированный порог · Отсу · K-means · SVM · U-Net',.85,2.3,6.5,1,25,TEAL,True)
slide('Задание','Пять методов сегментации','1. Фиксированный порог\n2. Отсу\n3. K-means, k=5\n4. SVM объект/фон\n5. Сравнение с U-Net',24)
slide('Теория','Что такое сегментация','Сегментация присваивает каждому пикселю метку области.\n\nГлобальный порог использует одно t. Отсу выбирает t автоматически. K-means группирует признаки. SVM строит границу классов. U-Net преобразует изображение в маску с помощью encoder-decoder и skip connections.',21)
slide('Фиксированный порог и Отсу','Бинаризация яркости','B(x,y)=255, если I(x,y) >= t, иначе 0.\n\nРучной порог: 180.\nПорог Отсу: 141.\n\nОтсу выбирает порог по максимуму межклассовой дисперсии и не требует ручного подбора.',22)
slide('K-means, k=5','Кластеризация пикселей','K-means чередует назначение пикселей ближайшему центру и пересчёт центров.\n\nДля изображения получены центры яркости: 12.18, 69.61, 132.42, 201.44, 227.33.\n\nКластеры не имеют заранее заданной семантики объекта и фона.',21)
slide('SVM для пиксельной сегментации','Признаки: яркость, x и y','Каждый пиксель рассматривается как объект классификации. RBF-SVM обучается разделять объект и фон.\n\nЦели взяты из маски Отсу, так как ручной ground truth отсутствует. Поэтому высокий IoU SVM отражает приближение к Отсу, а не независимую точность.',21)
slide('U-Net на PyTorch','Сегментационная нейронная сеть','Tiny U-Net: Conv -> Pool -> Conv -> Upsample -> Skip connection -> маска.\n\nОбучение: BCEWithLogitsLoss, Adam, L2 weight_decay=1e-4, 8 эпох.\n\nОбучающие патчи получают псевдоцели Отсу.',21)
slide('Результаты','IoU относительно маски Отсу','Фиксированный порог 180: 0.9505\nОтсу: 1.0000\nSVM: 0.9995\nU-Net: 0.8504\n\nU-Net обучалась коротко и на небольшом наборе патчей. Для честной оценки нужна независимая ручная маска.',23)
s=slide('Сравнение методов','Численные результаты'); s.shapes.add_picture(str(chart),Inches(.85),Inches(1.85),width=Inches(8),height=Inches(4.7)); text(s,'SVM почти повторяет Отсу, потому что обучался на его псевдоразметке.\n\nIoU: пересечение масок / их объединение.',9.2,2.5,3.1,2,19,TEAL,True)
slide('Код PyTorch','Фрагмент U-Net','class TinyUNet(nn.Module):\n    ...\n    def forward(self, x):\n        a = self.encoder(x)\n        b = self.pool(a)\n        return self.decoder(b, a)\n\nloss = BCEWithLogitsLoss()(logits, target)\nloss.backward()\noptimizer.step()',20)
slide('Контрольные вопросы','Ответы 1–5','1. Сегментация делит изображение на области или классы пикселей.\n2. Глобальная бинаризация использует один порог, адаптивная — локальные пороги.\n3. Watershed строит линии водораздела между областями рельефа яркости.\n4. K-means даёт жёсткий кластер, Fuzzy C-means — степени принадлежности.\n5. CNN автоматически извлекают признаки, учитывают контекст и сложные границы; U-Net сохраняет детали skip connections.',20)
slide('Итоги','Что выполнено','✓ Фиксированный порог и Отсу\n✓ K-means с k=5\n✓ SVM объект/фон\n✓ U-Net на PyTorch\n✓ IoU и сравнительная диаграмма\n✓ Теория и ответы на контрольные вопросы\n\nОграничение: без ground truth сравнение SVM и U-Net с Отсу является демонстрационным.',23)
prs.save(OUT); print(OUT)
