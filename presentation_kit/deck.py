"""Конструктор слайдов в минималистичном Material-стиле поверх python-pptx."""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import measure as M
from . import tokens as T

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def rgb(value):
    return RGBColor.from_string(value)


def _soft_shadow(shape, blur=9.0, dist=2.0, alpha=8):
    """Мягкая тень Material-карточки: python-pptx не умеет её задавать напрямую."""
    sp_pr = shape._element.spPr
    for existing in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(existing)
    xml = (
        f'<a:effectLst xmlns:a="{_A}">'
        f'<a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}"'
        ' dir="5400000" rotWithShape="0">'
        f'<a:srgbClr val="1B2A44"><a:alpha val="{int(alpha * 1000)}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    sp_pr.append(etree.fromstring(xml))


def _no_shadow(shape):
    sp_pr = shape._element.spPr
    for existing in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(existing)
    sp_pr.append(etree.fromstring(f'<a:effectLst xmlns:a="{_A}"/>'))


def _drop_theme_style(shape):
    """Убирает <p:style>: иначе фигура тянет из темы обводку, тень и цвет текста."""
    element = shape._element
    for style in element.findall(qn("p:style")):
        element.remove(style)


def _warn(kind, needed, given):
    if needed > given + 1e-6:
        print(f"  ! {kind}: нужно {needed:.2f}\" при {given:.2f}\"")


def _letter_spacing(run, points):
    run.font._rPr.set("spc", str(int(points * 100)))


class Deck:
    """Презентация с общим оформлением: шапка, подвал, карточки, метрики."""

    def __init__(self, footer_left):
        self.prs = Presentation()
        self.prs.slide_width = Inches(T.SLIDE_W)
        self.prs.slide_height = Inches(T.SLIDE_H)
        self.footer_left = footer_left
        self._audit = []     # что и где стоит — для проверки вылетов за поля
        self._in_footer = False

    def _register(self, x, y, w, h, what):
        if not self._in_footer:
            self._audit.append((len(self.prs.slides._sldIdLst), x, y, w, h, what))

    # --- низкоуровневые примитивы -------------------------------------------

    def _blank(self, background=T.SURFACE):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(background)
        return slide

    def text(self, slide, content, x, y, w, h, size=T.BODY, color=T.INK, bold=False,
             align=PP_ALIGN.LEFT, line_spacing=1.28, space_after=6, font=T.FONT,
             anchor=MSO_ANCHOR.TOP, spacing=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = anchor
        for index, line in enumerate(str(content).split("\n")):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.text = line
            para.alignment = align
            para.line_spacing = line_spacing
            para.space_after = Pt(space_after)
            para.font.name = font
            para.font.size = Pt(size)
            para.font.bold = bold
            para.font.color.rgb = rgb(color)
            if spacing is not None and para.runs:
                _letter_spacing(para.runs[0], spacing)
        real_h = M.text_height(content, w, size, line_spacing, space_after, font, bold)
        self._register(x, y, w, real_h, str(content).split("\n")[0][:46])
        return box

    def rect(self, slide, x, y, w, h, fill=T.CARD, radius=T.RADIUS, shadow=False,
             line=None, line_width=1.0):
        shape_kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shape.adjustments[0] = min(radius / min(w, h), 0.5)
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_width)
        shape.text_frame.word_wrap = True
        self._register(x, y, w, h, "плашка")
        _drop_theme_style(shape)
        if shadow:
            _soft_shadow(shape)
        else:
            _no_shadow(shape)
        return shape

    def picture(self, slide, path, x, y, w=None, h=None, shadow=True):
        """Картинка по ширине или высоте; PNG уже со скруглёнными углами."""
        kwargs = {}
        if w is not None:
            kwargs["width"] = Inches(w)
        if h is not None:
            kwargs["height"] = Inches(h)
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)
        if shadow:
            _soft_shadow(pic)
        pic.bottom_in = (pic.top + pic.height) / 914400
        pic.width_in = pic.width / 914400
        self._register(x, y, pic.width_in, pic.height / 914400, Path(path).name)
        return pic

    def fit_picture(self, slide, path, x, y, max_w, max_h, shadow=True, center=True):
        """Вписывает картинку в прямоугольник, сохраняя пропорции."""
        with Image.open(str(path)) as source:
            native_w, native_h = source.size
        aspect = native_h / native_w
        width = max_w if max_w * aspect <= max_h else max_h / aspect
        left = x + (max_w - width) / 2 if center else x
        return self.picture(slide, path, left, y, w=width, shadow=shadow)

    # --- шапка и подвал ------------------------------------------------------

    def _footer(self, slide, number):
        self._in_footer = True
        self.text(slide, self.footer_left, T.MARGIN, T.FOOTER_Y, 8, 0.24,
                  size=9, color=T.MUTED, space_after=0)
        self.text(slide, f"{number:02d}", T.SLIDE_W - T.MARGIN - 1.0, T.FOOTER_Y, 1.0, 0.24,
                  size=9, color=T.MUTED, align=PP_ALIGN.RIGHT, space_after=0)
        self._in_footer = False

    def slide(self, title, sub=None, kicker=None):
        slide = self._blank()
        if kicker:
            self.text(slide, kicker.upper(), T.MARGIN, T.KICKER_Y, T.CONTENT_W, 0.24,
                      size=T.KICKER, color=T.PRIMARY, bold=True, spacing=1.4, space_after=0)
        self.text(slide, title, T.MARGIN, T.TITLE_Y if kicker else T.TITLE_Y - 0.1,
                  T.CONTENT_W, 0.6, size=T.TITLE, bold=True, space_after=0, line_spacing=1.0)
        if sub:
            self.text(slide, sub, T.MARGIN, T.SUBTITLE_Y, T.CONTENT_W, 0.32,
                      size=T.SUBTITLE, color=T.MUTED, space_after=0)
        self._footer(slide, len(self.prs.slides._sldIdLst))
        slide.body_top = T.BODY_TOP if sub else T.BODY_TOP_BARE
        return slide

    # --- готовые блоки -------------------------------------------------------

    def cover(self, kicker, title, subtitle, meta, image=None, image_w=5.08,
              title_size=38, text_w=6.05):
        slide = self._blank()
        self._in_footer = True
        self.rect(slide, 0, 0, T.SLIDE_W, 0.10, fill=T.PRIMARY, radius=0)
        self._in_footer = False
        self.text(slide, kicker.upper(), T.MARGIN, 1.30, text_w, 0.26,
                  size=T.KICKER, color=T.PRIMARY, bold=True, spacing=1.6, space_after=0)
        self.text(slide, title, T.MARGIN, 1.80, text_w, 2.1,
                  size=title_size, bold=True, line_spacing=1.08, space_after=0)
        self.text(slide, subtitle, T.MARGIN, 4.24, text_w, 1.2,
                  size=T.BODY_LG, color=T.MUTED, line_spacing=1.34, space_after=0)
        self.text(slide, meta, T.MARGIN, 6.28, text_w + 1.5, 0.3,
                  size=T.BODY_SM, color=T.PRIMARY, bold=True, space_after=0)
        if image:
            pic = self.picture(slide, image, T.SLIDE_W - T.MARGIN - image_w, 0, w=image_w)
            height_in = pic.height / 914400
            pic.top = Emu(int(Inches((T.SLIDE_H - height_in) / 2 + 0.05)))
        self._footer(slide, len(self.prs.slides._sldIdLst))
        return slide

    def section(self, number, title, note=None):
        slide = self._blank(background=T.PRIMARY)
        self.text(slide, number, T.MARGIN, 2.30, 3.0, 1.4,
                  size=76, color="FFFFFF", bold=True, space_after=0, line_spacing=1.0)
        self.text(slide, title, T.MARGIN, 3.75, 9.4, 1.0,
                  size=30, color="FFFFFF", bold=True, space_after=0, line_spacing=1.1)
        if note:
            self.text(slide, note, T.MARGIN, 4.95, 8.6, 0.8,
                      size=T.BODY, color="C9DBFB", space_after=0)
        self.text(slide, f"{len(self.prs.slides._sldIdLst):02d}",
                  T.SLIDE_W - T.MARGIN - 1.0, T.FOOTER_Y, 1.0, 0.24,
                  size=9, color="9EBDF3", align=PP_ALIGN.RIGHT, space_after=0)
        return slide

    def steps(self, slide, items, x, y, w, size=T.BODY, row_h=None, chip=0.40, start=1):
        """Нумерованный список: квадратный чип с номером и текст рядом."""
        text_w = w - chip - 0.24
        needed = max(M.text_height(item, text_w, size, 1.2, 0) for item in items) + 0.30
        if row_h is None:
            row_h = max(needed, chip + 0.22)
        else:
            _warn("steps", needed, row_h)
        for index, item in enumerate(items):
            top = y + index * row_h
            box = self.rect(slide, x, top, chip, chip, fill=T.PRIMARY_CONTAINER,
                            radius=0.10)
            self.text(slide, f"{index + start:02d}", x, top + 0.075, chip, 0.24,
                      size=9.5, color=T.ON_PRIMARY_CONTAINER, bold=True,
                      align=PP_ALIGN.CENTER, space_after=0)
            self.text(slide, item, x + chip + 0.24, top + 0.02, w - chip - 0.24, row_h,
                      size=size, color=T.INK, space_after=0, line_spacing=1.2)
            del box
        return y + len(items) * row_h

    @staticmethod
    def _value_size(value, inner_w, size=T.METRIC, floor=13.0):
        """Уменьшает кегль значения, если оно не влезает в карточку одной строкой."""
        limit = inner_w * 72
        while size > floor and M.width_pt(str(value), size, bold=True) > limit:
            size -= 0.5
        return size

    def metrics(self, slide, items, x, y, w, h=1.12, gap=T.GAP, vertical=False,
                fill=T.CARD, value_color=T.PRIMARY):
        """Плитки «значение + подпись». items: [(значение, подпись), ...]"""
        count = len(items)
        if vertical:
            cell_h = (h - gap * (count - 1)) / count
            for index, (value, label) in enumerate(items):
                top = y + index * (cell_h + gap)
                self.rect(slide, x, top, w, cell_h, fill=fill, shadow=False)
                size = self._value_size(value, w - 0.56)
                block = 0.34 + 0.10 + M.text_height(label, w - 0.52, T.METRIC_LABEL, 1.2, 0)
                head = top + (cell_h - block) / 2
                self.text(slide, value, x + 0.28, head, w - 0.56, 0.40,
                          size=size, color=value_color, bold=True, space_after=0,
                          line_spacing=1.0)
                self.text(slide, label, x + 0.28, head + 0.44, w - 0.56, 0.30,
                          size=T.METRIC_LABEL, color=T.MUTED, space_after=0, line_spacing=1.18)
            return y + h
        cell_w = (w - gap * (count - 1)) / count
        for index, (value, label) in enumerate(items):
            left = x + index * (cell_w + gap)
            self.rect(slide, left, y, cell_w, h, fill=fill, shadow=False)
            size = self._value_size(value, cell_w - 0.56)
            block = 0.34 + 0.10 + M.text_height(label, cell_w - 0.56, T.METRIC_LABEL, 1.2, 0)
            head = y + (h - block) / 2
            self.text(slide, value, left + 0.28, head, cell_w - 0.56, 0.40,
                      size=size, color=value_color, bold=True, space_after=0,
                      line_spacing=1.0)
            self.text(slide, label, left + 0.28, head + 0.44, cell_w - 0.56, 0.30,
                      size=T.METRIC_LABEL, color=T.MUTED, space_after=0, line_spacing=1.18)
        return y + h

    def note(self, slide, content, x, y, w, h=None, size=T.BODY_SM, color=T.MUTED,
             fill=T.CARD, bar=False):
        """Карточка-пояснение; bar=True добавляет вертикальную акцентную линию."""
        pad = 0.60 if bar else 0.34
        inner = w - pad - 0.32
        needed = M.text_height(content, inner, size, 1.30, 5) + 0.56
        if h is None:
            h = needed
        else:
            _warn("note", needed, h)
        self.rect(slide, x, y, w, h, fill=fill, shadow=False)
        if bar:
            self.rect(slide, x + 0.26, y + 0.26, 0.055, h - 0.52, fill=T.PRIMARY, radius=0.03)
        self.text(slide, content, x + pad, y + (h - needed + 0.56) / 2, inner, h - 0.56,
                  size=size, color=color, space_after=5, line_spacing=1.30)
        return y + h

    def code(self, slide, content, x, y, w, h=None, size=T.CODE):
        needed = M.text_height(content, w - 0.72, size, 1.28, 1, font=T.FONT_MONO) + 0.52
        if h is None:
            h = needed
        else:
            _warn("code", needed, h)
        self.rect(slide, x, y, w, h, fill=T.CARD, shadow=False)
        self.rect(slide, x, y + 0.24, 0.055, h - 0.48, fill=T.PRIMARY, radius=0.03)
        self.text(slide, content, x + 0.42, y + 0.26, w - 0.72, h - 0.52,
                  size=size, color=T.INK, font=T.FONT_MONO, space_after=1,
                  line_spacing=1.28)
        return y + h

    def caption(self, slide, content, x, y, w, size=T.CAPTION):
        self.text(slide, content, x, y, w, 0.3, size=size, color=T.MUTED, space_after=0)

    def qa(self, slide, pairs, x, y, w, columns=2, gap=0.5, size=11.5, row_h=None,
           avail_h=None):
        """Компактные пары «вопрос — ответ» в одну или две колонки."""
        col_w = (w - gap * (columns - 1)) / columns
        per_col = -(-len(pairs) // columns)
        needed = max(M.text_height(q, col_w, size, 1.15, 0, bold=True)
                     + M.text_height(a, col_w, size - 0.5, 1.22, 0) + 0.36
                     for q, a in pairs)
        if row_h is None:
            row_h = max(needed, avail_h / per_col) if avail_h else needed
        else:
            _warn("qa", needed, row_h)
        for index, (question, answer) in enumerate(pairs):
            col = index // per_col
            row = index % per_col
            left = x + col * (col_w + gap)
            top = y + row * row_h
            q_h = M.text_height(question, col_w, size, 1.15, 0, bold=True)
            self.text(slide, question, left, top, col_w, q_h,
                      size=size, color=T.INK, bold=True, space_after=0, line_spacing=1.15)
            self.text(slide, answer, left, top + q_h + 0.07, col_w, row_h - q_h - 0.10,
                      size=size - 0.5, color=T.MUTED, space_after=0, line_spacing=1.22)

    def columns(self, slide, blocks, x, y, w, h=None, gap=T.GAP, fill=T.CARD, size=T.BODY_SM):
        """Ряд карточек «заголовок + текст»."""
        count = len(blocks)
        cell_w = (w - gap * (count - 1)) / count
        needed = max(M.text_height(content, cell_w - 0.60, size, 1.28, 4)
                     for _, content in blocks) + 1.02
        if h is None:
            h = needed
        else:
            _warn("columns", needed, h)
        for index, (head, content) in enumerate(blocks):
            left = x + index * (cell_w + gap)
            self.rect(slide, left, y, cell_w, h, fill=fill, shadow=False)
            self.text(slide, head, left + 0.30, y + 0.30, cell_w - 0.60, 0.3,
                      size=T.BODY_SM, color=T.PRIMARY, bold=True, space_after=0)
            self.text(slide, content, left + 0.30, y + 0.72, cell_w - 0.60, h - 1.0,
                      size=size, color=T.MUTED, space_after=4, line_spacing=1.28)
        return y + h

    def checklist(self, slide, items, x, y, w, size=T.BODY, row_h=None):
        needed = max(M.text_height(item, w - 0.42, size, 1.2, 0) for item in items) + 0.16
        if row_h is None:
            row_h = needed
        else:
            _warn("checklist", needed, row_h)
        for index, item in enumerate(items):
            top = y + index * row_h
            self.rect(slide, x, top + 0.10, 0.16, 0.16, fill=T.POSITIVE, radius=0.08)
            self.text(slide, item, x + 0.42, top, w - 0.42, row_h,
                      size=size, color=T.INK, space_after=0, line_spacing=1.2)
        return y + len(items) * row_h

    def table(self, slide, header, rows, x, y, w, size=11.5, row_h=0.40,
              widths=None, highlight=None):
        """Лёгкая таблица без рамок: только заголовок, линейки и текст."""
        count = len(header)
        widths = widths or [1.0 / count] * count
        xs, acc = [], 0.0
        for fraction in widths:
            xs.append(x + acc * w)
            acc += fraction
        for index, cell in enumerate(header):
            self.text(slide, cell, xs[index], y, widths[index] * w - 0.12, 0.26,
                      size=size - 1, color=T.MUTED, bold=True, space_after=0)
        self.rect(slide, x, y + 0.30, w, 0.014, fill=T.OUTLINE, radius=0)
        for row_index, row in enumerate(rows):
            top = y + 0.40 + row_index * row_h
            accent = highlight is not None and row_index == highlight
            if accent:
                self.rect(slide, x - 0.14, top - 0.06, w + 0.28, row_h - 0.02,
                          fill=T.PRIMARY_CONTAINER, radius=0.07)
            for index, cell in enumerate(row):
                self.text(slide, cell, xs[index], top, widths[index] * w - 0.12, 0.26,
                          size=size, color=T.INK if accent else T.INK,
                          bold=accent, space_after=0)
            if row_index < len(rows) - 1:
                self.rect(slide, x, top + row_h - 0.08, w, 0.008, fill=T.OUTLINE, radius=0)
        return y + 0.40 + len(rows) * row_h

    def report(self, bottom_limit=T.FOOTER_Y - 0.12, right_limit=T.SLIDE_W - 0.55):
        """Печатает всё, что вылезает за нижнее поле или за правый край."""
        problems = []
        for number, x, y, w, h, what in self._audit:
            if y + h > bottom_limit + 1e-3:
                problems.append(f"  слайд {number:02d}: «{what}» ниже поля на "
                                f"{(y + h - bottom_limit):.2f}\"")
            if x + w > right_limit + 1e-3:
                problems.append(f"  слайд {number:02d}: «{what}» правее поля на "
                                f"{(x + w - right_limit):.2f}\"")
        for line in problems:
            print(line)
        return problems

    def save(self, path):
        path = Path(path)
        self.report()
        self.prs.save(str(path))
        return path, len(self.prs.slides._sldIdLst)
